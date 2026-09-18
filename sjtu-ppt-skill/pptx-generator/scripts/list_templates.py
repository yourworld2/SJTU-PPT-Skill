#!/usr/bin/env python3
"""
Template Indexer for pptx-generator.
Scan the assets directory, analyze all .pptx templates, and output an index
with layout mappings and placeholder information for each template.
"""

import argparse
import json
import os
import sys

from pptx import Presentation


def classify_layout(layout_name):
    """Classify a layout name into a canonical type."""
    name_lower = layout_name.lower()
    if any(k in name_lower for k in ['封面', 'cover', 'title']):
        return 'cover'
    elif any(k in name_lower for k in ['目录', 'toc']):
        return 'toc'
    elif any(k in name_lower for k in ['过渡', 'section', 'divider', 'chapter', '自定义']):
        return 'section'
    elif any(k in name_lower for k in ['正文', '内容', 'body', 'text', '内页']):
        return 'content'
    elif any(k in name_lower for k in ['封底', 'end', 'thank', 'closing']):
        return 'end'
    return 'unknown'


def analyze_template(template_path):
    """Analyze a single template and return its metadata."""
    prs = Presentation(template_path)

    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        layout_type = classify_layout(layout.name)
        placeholders = []
        for ph in layout.placeholders:
            ph_type = str(ph.placeholder_format.type)
            placeholders.append({
                'idx': ph.placeholder_format.idx,
                'type': ph_type,
                'name': ph.name
            })
        layouts.append({
            'index': i,
            'name': layout.name,
            'type': layout_type,
            'placeholders': placeholders
        })

    return {
        'filename': os.path.basename(template_path),
        'slide_width': prs.slide_width,
        'slide_height': prs.slide_height,
        'layout_count': len(prs.slide_layouts),
        'layouts': layouts
    }


def main():
    parser = argparse.ArgumentParser(description='Index all PPTX templates in assets directory.')
    parser.add_argument('--assets-dir', default='assets',
                        help='Directory containing .pptx template files (default: assets)')
    args = parser.parse_args()

    if not os.path.isdir(args.assets_dir):
        result = {
            'status': 'error',
            'message': f'Assets directory not found: {args.assets_dir}'
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        sys.exit(1)

    templates = []
    for filename in sorted(os.listdir(args.assets_dir)):
        if filename.lower().endswith('.pptx'):
            path = os.path.join(args.assets_dir, filename)
            try:
                info = analyze_template(path)
                templates.append(info)
            except Exception as e:
                templates.append({
                    'filename': filename,
                    'error': str(e)
                })

    result = {
        'status': 'success',
        'assets_dir': args.assets_dir,
        'template_count': len(templates),
        'templates': templates
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
